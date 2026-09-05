import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

NAVY = RGBColor(27, 54, 93)       # #1B365D
COBALT = RGBColor(43, 84, 126)    # #2B547E
SLATE = RGBColor(92, 118, 141)    # #5C768D
CHARCOAL = RGBColor(40, 40, 40)   # #282828

HEX_NAVY = "1B365D"
HEX_LIGHT_BG = "F4F6F9"
HEX_FORMULA_BG = "F8F9FA"
HEX_CALLOUT_BG = "EEF3F8"
HEX_CODE_BG = "F5F5F5"

def set_cell_bg(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_pads(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_tbl_borders(table, color="D0D7DE", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_runs_with_formatting(p, text, font_size=10.5, color=CHARCOAL):
    parts = re.split(r'(\*\*.*?\*\*|`.*?`|\$.*?\$)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2])
            r.font.name = 'Calibri'
            r.font.size = Pt(font_size)
            r.font.bold = True
            r.font.color.rgb = color
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(font_size * 0.92)
            r.font.color.rgb = COBALT
        elif part.startswith('$') and part.endswith('$'):
            r = p.add_run(clean_math(part[1:-1]))
            r.font.name = 'Cambria Math'
            r.font.size = Pt(font_size)
            r.font.italic = True
            r.font.color.rgb = NAVY
        else:
            r = p.add_run(part)
            r.font.name = 'Calibri'
            r.font.size = Pt(font_size)
            r.font.color.rgb = color

def clean_math(txt):
    repl = [
        ('\\text{', ''), ('\\mathbf{', ''), ('\\hat{', '^'), ('\\cdot', '·'),
        ('\\times', '×'), ('\\approx', '≈'), ('\\Delta', 'Δ'), ('\\pi', 'π'),
        ('\\sigma', 'σ'), ('\\lambda', 'λ'), ('\\psi', 'ψ'), ('\\theta', 'θ'),
        ('\\mathbb{I}', 'I'), ('\\in', '∈'), ('\\le', '≤'), ('\\ge', '≥'),
        ('\\sqrt{', '√('), ('\\frac{', '('), ('}{', ')/('), ('}', ')')
    ]
    for old, new in repl:
        txt = txt.replace(old, new)
    return txt.replace('{', '').replace('}', '')

def add_callout_box(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_bg(cell, HEX_CALLOUT_BG)
    set_cell_pads(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    rt = p.add_run(f"📌 {title}\n")
    rt.font.name = 'Calibri'
    rt.font.size = Pt(10)
    rt.font.bold = True
    rt.font.color.rgb = NAVY
    add_runs_with_formatting(p, text.strip(), font_size=9.5, color=CHARCOAL)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_formula_box(doc, eq_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_bg(cell, HEX_FORMULA_BG)
    set_cell_pads(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="2B547E"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(clean_math(eq_text.strip()))
    r.font.name = 'Cambria Math'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = NAVY
    
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_bg(cell, HEX_CODE_BG)
    set_cell_pads(cell, top=100, bottom=100, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(code_text.strip())
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = CHARCOAL
    
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        h = s.header.paragraphs[0]
        h.text = "CAMBODIA CPI PIPELINE  |  INFLATION NOWCASTING METHODOLOGY GUIDE"
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h.runs[0].font.name = "Calibri"
        h.runs[0].font.size = Pt(8.5)
        h.runs[0].font.color.rgb = RGBColor(120, 120, 120)
        
        f_p = s.footer.paragraphs[0]
        f_p.text = "CONFIDENTIAL & PROPRIETARY  |  AUTOMATED RETAIL PRICE PLATFORM"
        f_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        f_p.runs[0].font.name = "Calibri"
        f_p.runs[0].font.size = Pt(8.5)
        f_p.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    # Document Header
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(12)
    p_t.paragraph_format.space_after = Pt(4)
    rt = p_t.add_run("High-Frequency Inflation Nowcasting:\nEconometric Foundations, Mathematical Formulations, and Implementation Guide")
    rt.font.name = "Calibri"
    rt.font.size = Pt(21)
    rt.font.bold = True
    rt.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    rs = p_sub.add_run("Production Methodology for the Automated Cambodia Consumer Price Index (CPI) Platform\nVersion 2.4 — Post-Audit Verified")
    rs.font.name = "Calibri"
    rs.font.size = Pt(11.5)
    rs.font.italic = True
    rs.font.color.rgb = SLATE

    i = 0
    in_code = False
    code_buf = []
    in_eq = False
    eq_buf = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i]
        s_line = line.strip()

        # Code block handler
        if s_line.startswith('```'):
            if in_code:
                add_code_block(doc, '\n'.join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line.rstrip())
            i += 1
            continue

        # Multiline equation handler ($$...$$)
        if s_line.startswith('$$'):
            if in_eq:
                eq_buf.append(s_line[2:] if s_line.endswith('$$') and len(s_line) > 2 else '')
                add_formula_box(doc, '\n'.join(eq_buf))
                eq_buf = []
                in_eq = False
            else:
                if s_line.endswith('$$') and len(s_line) > 2:
                    add_formula_box(doc, s_line[2:-2])
                else:
                    in_eq = True
                    eq_buf.append(s_line[2:])
            i += 1
            continue

        if in_eq:
            if s_line.endswith('$$'):
                eq_buf.append(s_line[:-2])
                add_formula_box(doc, '\n'.join(eq_buf))
                eq_buf = []
                in_eq = False
            else:
                eq_buf.append(s_line)
            i += 1
            continue

        # Table handler
        if s_line.startswith('|') and s_line.endswith('|'):
            if not in_table:
                in_table = True
                table_rows = []
            if not re.match(r'\|(\s*:?-+:?\s*\|)+', s_line):
                cells = [c.strip() for c in s_line.split('|')[1:-1]]
                table_rows.append(cells)
            i += 1
            continue
        elif in_table:
            # flush table
            if table_rows:
                tbl = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = False
                set_tbl_borders(tbl)
                for r_idx, r_cells in enumerate(table_rows):
                    bg = HEX_NAVY if r_idx == 0 else (HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF")
                    for c_idx, val in enumerate(r_cells):
                        c = tbl.cell(r_idx, c_idx)
                        set_cell_bg(c, bg)
                        set_cell_pads(c, top=60, bottom=60, left=80, right=80)
                        p_c = c.paragraphs[0]
                        if r_idx == 0:
                            r = p_c.add_run(val)
                            r.font.name = "Calibri"
                            r.font.size = Pt(9.5)
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            add_runs_with_formatting(p_c, val, font_size=9, color=CHARCOAL)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            in_table = False
            table_rows = []

        # Ignore horizontal rules
        if s_line.startswith('---'):
            i += 1
            continue

        # Skip main title (already rendered)
        if s_line.startswith('# ') and 'High-Frequency' in s_line:
            i += 1
            continue

        # Headings
        if s_line.startswith('# '):
            h = doc.add_heading(s_line[2:].strip(), level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(5)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(15)
                r.font.bold = True
                r.font.color.rgb = NAVY
            i += 1
            continue

        if s_line.startswith('## '):
            h = doc.add_heading(s_line[3:].strip(), level=2)
            h.paragraph_format.space_before = Pt(13)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(12.5)
                r.font.bold = True
                r.font.color.rgb = COBALT
            i += 1
            continue

        if s_line.startswith('### '):
            h = doc.add_heading(s_line[4:].strip(), level=3)
            h.paragraph_format.space_before = Pt(9)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = CHARCOAL
            i += 1
            continue

        # Blockquote handler
        if s_line.startswith('> '):
            add_callout_box(doc, s_line[2:].strip(), title="KEY TAKEAWAY / PRINCIPLE")
            i += 1
            continue

        # Bullet lists
        if s_line.startswith('* ') or s_line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_runs_with_formatting(p, s_line[2:].strip(), font_size=10, color=CHARCOAL)
            i += 1
            continue

        # Numbered lists
        if re.match(r'^\d+\.\s', s_line):
            num_match = re.match(r'^(\d+\.)\s*(.*)', s_line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(num_match.group(1) + " ")
            r_num.font.name = 'Calibri'
            r_num.font.size = Pt(10)
            r_num.font.bold = True
            r_num.font.color.rgb = NAVY
            add_runs_with_formatting(p, num_match.group(2).strip(), font_size=10, color=CHARCOAL)
            i += 1
            continue

        # Standard Paragraph
        if s_line:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            add_runs_with_formatting(p, s_line, font_size=10, color=CHARCOAL)

        i += 1

    doc.save(docx_path)
    print(f"SUCCESS: Generated {docx_path}")

if __name__ == '__main__':
    md = os.path.abspath('docs/NOWCASTING_METHODOLOGY_GUIDE.md')
    out = os.path.abspath('docs/CAMBODIA_CPI_NOWCASTING_METHODOLOGY_GUIDE.docx')
    convert_md_to_docx(md, out)
